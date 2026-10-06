"""Tests for the GSTT ACR protocol."""

# ruff: noqa: PT009 PT027 SLF001

# Python imports
import io
import json
import os
import shutil
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from typing import Any
from unittest.mock import Mock, patch

# Module imports
import openpyxl
import pydicom
import pytest
from PIL import Image as PILImage

# Local imports
import hazenlib
from hazenlib.orchestration import init_task
from hazenlib.protocols.acr_gstt import (
    ROLES,
    WORKBOOK_TEMPLATE,
    ACRGSTTProtocol,
    Session,
    Step,
    StepOutcome,
    build_steps,
    detect_vendor,
    match_role,
    rows_as_records,
)
from hazenlib.protocols.outputs import (
    ImageGroup,
    TemplateCell,
    fill_template,
    write_records,
)
from hazenlib.types import Measurement, Result
from tests import TEST_DATA_DIR

SIEMENS_T1 = TEST_DATA_DIR / "acr" / "Siemens_Sola_1.5T_T1"
GE_T1 = TEST_DATA_DIR / "acr" / "GE_Artist_1.5T_T1"
PHILIPS_SESSION = (
    TEST_DATA_DIR / "acr" / "Philips_Ingenia_Ambition_1.5T_all_series"
)


def make_session(root: Path, roles: dict[str, Path]) -> Path:
    """Create a session folder of symlinked series named by role."""
    root.mkdir(parents=True, exist_ok=True)
    for name, src in roles.items():
        (root / name).symlink_to(src, target_is_directory=True)
    return root


def fake_result(task: str, *measurements: Measurement) -> Result:
    """Build a Result without reading DICOM metadata."""
    result = Result(task=task, _load_metadata=False)
    for m in measurements:
        result.add_measurement(m)
    result.add_report_image(f"/report/{task}.png")
    return result


SLICE_WIDTH = Measurement("SliceWidth", 4.6, subtype="slice width")
UNIFORMITY = Measurement("Uniformity", 97.0, subtype="Integral uniformity")
SNR = Measurement("SNR", 2000.0, type="normalised", subtype="subtraction")


def fake_init_task(task_name: str, files: list, **kwargs) -> Mock:
    """Stand-in for init_task returning canned results per task."""
    measurement = {
        "acr_slice_thickness": SLICE_WIDTH,
        "acr_uniformity": UNIFORMITY,
        "acr_snr": SNR,
    }.get(task_name, UNIFORMITY)
    task = Mock()
    task.run.return_value = fake_result(task_name, measurement)
    return task


class TestRoleMatching(unittest.TestCase):
    """Folder names map to roles regardless of case, separators and order."""

    def test_matches(self) -> None:
        for name, role in {
            "Head_1": "Head_1",
            "HEAD1": "Head_1",
            "head-5": "Head_5",
            "Body_Sag_4": "Body_Sag_4",
            "body sag 4": "Body_Sag_4",
            "BODY_TRA_1": "Body_Tra_1",
            "Body_Ax_1": "Body_Tra_1",
            "BODY1_ax": "Body_Tra_1",
            "BODY2_cor": "Body_Cor_2",
            "body3 axial": "Body_Tra_3",
            "BODY4_sag_rotMAT": "Body_Sag_4",
            "Body_Sag_4_rotated": "Body_Sag_4",
            "HEAD2_rep": "Head_2",
            "Head1_ax": "Head_1",
        }.items():
            with self.subTest(name=name):
                self.assertEqual(match_role(name), role)

    def test_non_matches(self) -> None:
        for name in (
            "RF_noise1",
            "Head_6",
            "Head_12",
            "Body_Sag_5",
            "BODY1",  # no plane
            "Body_Sag_1_cor",  # two planes
            "Head_Sag_1",
            "Body1axx",
            "my_Head_1",
        ):
            with self.subTest(name=name):
                self.assertIsNone(match_role(name))

    def test_role_count(self) -> None:
        self.assertEqual(len(ROLES), 17)


class TestBuildSteps(unittest.TestCase):
    """Vendor rules from the GSTT work instruction."""

    @staticmethod
    def roles_for(steps: tuple[Step, ...], test: str) -> list[str]:
        return [s.role for s in steps if s.test == test]

    def test_siemens(self) -> None:
        steps = build_steps("siemens")
        self.assertEqual(len(steps), 21)
        self.assertEqual(self.roles_for(steps, "ghosting"), ["Head_1"])
        self.assertEqual(
            self.roles_for(steps, "geometric_accuracy"),
            ["Head_4", "Body_Tra_4", "Body_Sag_4", "Body_Cor_4"],
        )
        self.assertEqual(
            self.roles_for(steps, "uniformity"),
            ["Head_3", "Body_Tra_3", "Body_Sag_3", "Body_Cor_3"],
        )

    def test_philips_ghosting_uses_head_5(self) -> None:
        steps = build_steps("philips")
        self.assertEqual(self.roles_for(steps, "ghosting"), ["Head_5"])

    def test_ge_uses_series_1(self) -> None:
        steps = build_steps("ge")
        for test in (
            "slice_thickness",
            "slice_position",
            "geometric_accuracy",
        ):
            with self.subTest(test=test):
                self.assertEqual(
                    self.roles_for(steps, test),
                    ["Head_1", "Body_Tra_1", "Body_Sag_1", "Body_Cor_1"],
                )
        self.assertEqual(
            self.roles_for(steps, "uniformity"),
            ["Head_3", "Body_Tra_1", "Body_Sag_1", "Body_Cor_1"],
        )

    def test_snr_pairs_and_slice_width_source(self) -> None:
        snr = [s for s in build_steps("siemens") if s.test == "snr"]
        self.assertEqual(
            [(s.role, s.subtract, s.slice_width_from) for s in snr],
            [
                ("Head_1", "Head_2", "Head_4"),
                ("Body_Tra_1", "Body_Tra_2", "Body_Tra_4"),
                ("Body_Sag_1", "Body_Sag_2", "Body_Sag_4"),
                ("Body_Cor_1", "Body_Cor_2", "Body_Cor_4"),
            ],
        )


class TestSessionAndVendor(unittest.TestCase):
    """Session discovery and vendor detection on real DICOM folders."""

    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp())

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp)

    def test_roles_found_and_others_ignored(self) -> None:
        root = make_session(
            self.tmp / "s",
            {"HEAD1": SIEMENS_T1, "Body_Sag_4": SIEMENS_T1},
        )
        (root / "RF_noise1").mkdir()
        (root / "Head_2").mkdir()  # empty, so ignored
        (root / "hazen_report").mkdir()  # earlier output, not listed
        session = Session.from_dir(root)
        self.assertEqual(list(session.folders), ["Head_1", "Body_Sag_4"])
        self.assertEqual(
            sorted(p.name for p in session.ignored),
            ["Head_2", "RF_noise1"],
        )

    def test_duplicate_role_raises(self) -> None:
        root = make_session(
            self.tmp / "s",
            {"Head_1": SIEMENS_T1, "head-1": SIEMENS_T1},
        )
        with self.assertRaisesRegex(ValueError, "both match role Head_1"):
            Session.from_dir(root)

    def test_missing_session_raises(self) -> None:
        with self.assertRaises(NotADirectoryError):
            Session.from_dir(self.tmp / "nope")

    def test_vendor_detected(self) -> None:
        root = make_session(self.tmp / "s", {"Head_1": SIEMENS_T1})
        self.assertEqual(detect_vendor(Session.from_dir(root)), "siemens")

    def test_mixed_vendors_raise(self) -> None:
        root = make_session(
            self.tmp / "s",
            {"Head_1": SIEMENS_T1, "Head_2": GE_T1},
        )
        with self.assertRaisesRegex(ValueError, "single scanner"):
            detect_vendor(Session.from_dir(root))

    def test_override_wins(self) -> None:
        root = make_session(
            self.tmp / "s",
            {"Head_1": SIEMENS_T1, "Head_2": GE_T1},
        )
        self.assertEqual(
            detect_vendor(Session.from_dir(root), override="Philips"),
            "philips",
        )

    def test_canon_refused(self) -> None:
        canon = self.tmp / "canon"
        canon.mkdir()
        for src in sorted(SIEMENS_T1.iterdir()):
            ds = pydicom.dcmread(src)
            ds.Manufacturer = "Canon Medical Systems"
            ds.save_as(canon / src.name)
        root = make_session(self.tmp / "s", {"Head_1": canon})
        with self.assertRaisesRegex(ValueError, "'canon' is not covered"):
            detect_vendor(Session.from_dir(root))

    def test_no_roles_raises(self) -> None:
        (self.tmp / "s" / "RF_noise1").mkdir(parents=True)
        with self.assertRaisesRegex(ValueError, "No role folders"):
            ACRGSTTProtocol(self.tmp / "s")


@patch("hazenlib.protocols.acr_gstt.init_task", side_effect=fake_init_task)
class TestRun(unittest.TestCase):
    """Staging, skipping and failure handling, with tasks mocked out."""

    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp())
        self.root = make_session(
            self.tmp / "s",
            {f"Head_{n}": SIEMENS_T1 for n in range(1, 6)},
        )

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp)

    def run_protocol(self) -> tuple[ACRGSTTProtocol, list]:
        protocol = ACRGSTTProtocol(self.root, report=False)
        return protocol, protocol.run(debug=True)

    @staticmethod
    def outcome(outcomes: list, role: str, test: str):  # noqa: ANN205
        return next(
            o for o in outcomes if o.step.role == role and o.step.test == test
        )

    def test_missing_roles_are_skipped(self, _init: Mock) -> None:
        _, outcomes = self.run_protocol()
        head = [o for o in outcomes if o.step.role.startswith("Head")]
        body = [o for o in outcomes if o.step.role.startswith("Body")]
        self.assertTrue(all(o.status == "ok" for o in head))
        self.assertTrue(all(o.status == "skipped" for o in body))
        self.assertEqual(
            self.outcome(outcomes, "Body_Tra_3", "uniformity").detail,
            "missing Body_Tra_3",
        )

    def test_snr_uses_measured_slice_width(self, init: Mock) -> None:
        _, outcomes = self.run_protocol()
        snr_call = next(c for c in init.call_args_list if c[0][0] == "acr_snr")
        self.assertEqual(snr_call.kwargs["measured_slice_width"], 4.6)
        self.assertEqual(
            Path(snr_call.kwargs["subtract"]),
            self.root.resolve() / "Head_2",
        )
        snr = self.outcome(outcomes, "Head_1", "snr")
        self.assertEqual(snr.slice_width, 4.6)
        self.assertEqual(snr.slice_width_source, "Head_4")

    def test_snr_falls_back_to_nominal(self, init: Mock) -> None:
        def broken_thickness(task_name: str, files: list, **kwargs) -> Mock:
            if task_name == "acr_slice_thickness":
                msg = "no ramps found"
                raise RuntimeError(msg)
            return fake_init_task(task_name, files, **kwargs)

        init.side_effect = broken_thickness
        protocol, outcomes = self.run_protocol()

        thickness = self.outcome(outcomes, "Head_4", "slice_thickness")
        self.assertEqual(thickness.status, "failed")
        self.assertIn("RuntimeError: no ramps found", thickness.detail)

        snr_call = next(c for c in init.call_args_list if c[0][0] == "acr_snr")
        self.assertIsNone(snr_call.kwargs["measured_slice_width"])
        snr = self.outcome(outcomes, "Head_1", "snr")
        self.assertEqual(snr.status, "ok")
        self.assertTrue(snr.slice_width_source.startswith("nominal"))

        width_row = next(
            r
            for r in protocol.rows(outcomes)
            if r.role == "Head_1" and r.measurement == "SliceWidth"
        )
        self.assertEqual(width_row.type, "nominal")

    def test_snr_without_repeat_is_skipped(self, init: Mock) -> None:
        (self.root / "Head_2").unlink()
        _, outcomes = self.run_protocol()
        snr = self.outcome(outcomes, "Head_1", "snr")
        self.assertEqual(snr.status, "skipped")
        self.assertEqual(snr.detail, "missing Head_2")
        self.assertNotIn(
            "acr_snr",
            [c[0][0] for c in init.call_args_list],
        )

    def test_system_exit_is_a_failure(self, init: Mock) -> None:
        def exits(task_name: str, files: list, **kwargs) -> Mock:
            if task_name == "acr_ghosting":
                raise SystemExit
            return fake_init_task(task_name, files, **kwargs)

        init.side_effect = exits
        _, outcomes = self.run_protocol()
        self.assertEqual(
            self.outcome(outcomes, "Head_1", "ghosting").status,
            "failed",
        )

    def test_report_dir_per_role(self, init: Mock) -> None:
        protocol = ACRGSTTProtocol(self.root, report_dir=self.tmp / "rep")
        protocol.run(debug=True)
        uniformity = next(
            c for c in init.call_args_list if c[0][0] == "acr_uniformity"
        )
        self.assertEqual(
            Path(uniformity.kwargs["report_dir"]),
            self.tmp / "rep" / "Head_3",
        )

    def test_rows_are_labelled(self, _init: Mock) -> None:
        protocol, outcomes = self.run_protocol()
        rows = protocol.rows(outcomes)
        uniformity = next(r for r in rows if r.test == "uniformity")
        self.assertEqual(
            (uniformity.coil, uniformity.plane, uniformity.role),
            ("head", "tra", "Head_3"),
        )
        self.assertEqual(uniformity.value, 97.0)
        self.assertEqual(
            uniformity.report_images, "/report/acr_uniformity.png"
        )
        skipped = next(r for r in rows if r.role == "Body_Cor_3")
        self.assertEqual(skipped.status, "skipped")
        self.assertIsNone(skipped.value)

    def test_image_groups(self, _init: Mock) -> None:
        protocol = ACRGSTTProtocol(self.root, report_dir=self.tmp / "rep")
        sheets = protocol.image_groups(protocol.run(debug=True))
        self.assertEqual(
            list(sheets),
            [
                "Head Images",
                "Body Tra Images",
                "Body Sag Images",
                "Body Cor Images",
            ],
        )
        uniformity = next(
            g for g in sheets["Head Images"] if g.title == "Head_3 uniformity"
        )
        self.assertEqual(uniformity.paths, ("/report/acr_uniformity.png",))
        self.assertEqual(
            sheets["Body Cor Images"][0].note,
            "skipped: missing Body_Cor_4",
        )

    def test_no_image_groups_without_report(self, _init: Mock) -> None:
        protocol, outcomes = self.run_protocol()
        sheets = protocol.image_groups(outcomes)
        self.assertTrue(all(groups == [] for groups in sheets.values()))

    def test_workbook_cells(self, _init: Mock) -> None:
        protocol, outcomes = self.run_protocol()
        cells = {
            (c.sheet, c.ref): c for c in protocol.workbook_cells(outcomes)
        }
        # Head uniformity and slice thickness, normalised SNR
        self.assertEqual(cells["Hazen Output", "C16"].value, 97.0)
        self.assertEqual(cells["Hazen Output", "C23"].value, 4.6)
        self.assertEqual(cells["Hazen Output", "C15"].value, 2000.0)
        # Ran, but the canned result has no measured SNR
        self.assertEqual(
            cells["Hazen Output", "C14"].note,
            "Head_1 snr: not measured",
        )
        # Body roles are missing
        self.assertIsNone(cells["Hazen Output", "E16"].value)
        self.assertEqual(
            cells["Hazen Output", "E16"].note,
            "Body_Sag_3 uniformity skipped: missing Body_Sag_3",
        )
        geometry = [
            ref for sheet, ref in cells if sheet == "Geometric accuracy"
        ]
        self.assertEqual(
            geometry,
            [
                f"B{start + offset}"
                for start in (12, 26, 40, 54)
                for offset in (0, 1, 2, 3, 4, 5, 7)
            ],
        )


class TestOutputs(unittest.TestCase):
    """Workbook and record writers."""

    records = [
        {"role": "Head_3", "value": 97.0, "unit": "%"},
        {"role": "Head_5", "value": 0.08, "unit": "%"},
    ]

    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp())

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp)

    def test_fill_template(self) -> None:
        png = self.tmp / "img.png"
        PILImage.new("L", (800, 1600)).save(png)
        path = fill_template(
            WORKBOOK_TEMPLATE,
            [
                TemplateCell("Hazen Output", "C16", 97.0),
                TemplateCell("Geometric accuracy", "B12", note="skipped"),
            ],
            self.tmp / "out" / "r.xlsx",
            images={
                "Head Images": [
                    ImageGroup("Head_1 snr", (str(png), str(png))),
                    ImageGroup("Head_3 uniformity", note="skipped"),
                ],
                "Body Tra Images": [],
            },
        )
        wb = openpyxl.load_workbook(path)
        self.assertEqual(
            wb.sheetnames,
            [
                "Hazen Output",
                "Geometric accuracy",
                "Head Images",
                "Body Tra Images",
            ],
        )
        cover = wb["Hazen Output"]
        self.assertEqual(cover["C16"].value, 97.0)
        self.assertEqual(cover["B16"].value, "Uniformity (%)")
        header = cover.oddHeader
        self.assertTrue(header is None or header.left.text is None)
        self.assertIsNone(cover["A27"].value)
        head = wb["Head Images"]
        self.assertEqual(head["A1"].value, "Head_1 snr")
        with zipfile.ZipFile(path) as zf:
            media = [n for n in zf.namelist() if n.startswith("xl/media/")]
        self.assertEqual(len(media), 2)
        # 1600 px scaled to 720 px is 36 rows of 20 px, plus a gap
        self.assertEqual(head["A39"].value, "Head_3 uniformity")
        self.assertEqual(head["B39"].value, "skipped")
        self.assertEqual(wb["Body Tra Images"]["A1"].value, "No report images")
        geometry = wb["Geometric accuracy"]
        self.assertIsNone(geometry["B12"].value)
        self.assertEqual(geometry["B12"].comment.text, "skipped")

    def test_csv_and_json(self) -> None:
        write_records(self.records, "csv", self.tmp / "r.csv")
        self.assertEqual(
            (self.tmp / "r.csv").read_text().splitlines(),
            ["role,value,unit", "Head_3,97.0,%", "Head_5,0.08,%"],
        )
        write_records(self.records, "json", self.tmp / "r.json")
        self.assertEqual(
            json.loads((self.tmp / "r.json").read_text()),
            self.records,
        )

    def test_bad_format(self) -> None:
        with self.assertRaises(ValueError):
            write_records(self.records, "xml")


@pytest.mark.slow
class TestPhilipsSession(unittest.TestCase):
    """End to end on the scrubbed Philips Head 1-5 / Body Sag 1-4 data."""

    protocol: ACRGSTTProtocol
    outcomes: list[StepOutcome]

    @classmethod
    def setUpClass(cls) -> None:
        cls.protocol = ACRGSTTProtocol(PHILIPS_SESSION, report=False)
        cls.outcomes = cls.protocol.run()

    def test_statuses(self) -> None:
        self.assertEqual(self.protocol.vendor, "philips")
        ran = {o.step.role for o in self.outcomes if o.status == "ok"}
        self.assertEqual(
            ran,
            {"Head_1", "Head_3", "Head_4", "Head_5"}
            | {"Body_Sag_1", "Body_Sag_3", "Body_Sag_4"},
        )
        self.assertFalse(any(o.status == "failed" for o in self.outcomes))
        skipped = {o.step.role for o in self.outcomes if o.status == "skipped"}
        self.assertTrue(
            all(r.startswith(("Body_Tra", "Body_Cor")) for r in skipped)
        )

    def test_matches_individual_tasks(self) -> None:
        """Protocol output must equal running each task on its own."""
        for o in self.outcomes:
            if o.status != "ok":
                continue
            if o.result is None:
                self.fail(f"{o.step.role} {o.step.test} has no result")
            with self.subTest(role=o.step.role, test=o.step.test):
                kwargs: dict[str, Any] = {}
                if o.step.subtract:
                    kwargs["subtract"] = str(PHILIPS_SESSION / o.step.subtract)
                    kwargs["measured_slice_width"] = o.slice_width
                files = [
                    str(p) for p in (PHILIPS_SESSION / o.step.role).iterdir()
                ]
                expected = init_task(
                    o.step.task, files, report=False, **kwargs
                ).run()
                self.assertEqual(
                    o.result.measurements,
                    expected.measurements,
                )

    def test_snr_slice_width_from_head_4(self) -> None:
        snr = next(
            o
            for o in self.outcomes
            if o.step.role == "Head_1" and o.step.test == "snr"
        )
        thickness = next(
            o
            for o in self.outcomes
            if o.step.role == "Head_4" and o.step.test == "slice_thickness"
        )
        if thickness.result is None:
            self.fail("Head_4 slice thickness has no result")
        self.assertEqual(
            snr.slice_width,
            thickness.result.get_measurement(subtype="slice width")[0].value,
        )


@pytest.mark.slow
class TestCli(unittest.TestCase):
    """``hazen acr_gstt`` writes the workbook and returns 0."""

    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp())

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp)

    def test_cli(self) -> None:
        xlsx = self.tmp / "results.xlsx"
        argv = [
            "hazen",
            "acr_gstt",
            str(PHILIPS_SESSION),
            "--no-report",
            "--xlsx",
            str(xlsx),
            "--result",
            str(self.tmp / "results.csv"),
        ]
        with (
            patch.object(sys, "argv", argv),
            patch(
                "sys.stdout",
                new_callable=io.StringIO,
            ) as out,
        ):
            hazenlib.main()
        self.assertIn("failed 0", out.getvalue())
        wb = openpyxl.load_workbook(xlsx)
        self.assertEqual(
            wb.sheetnames[:3],
            ["Hazen Output", "Geometric accuracy", "Head Images"],
        )
        # Head uniformity from Head_3, body sagittal geometry from Body_Sag_4
        self.assertIsInstance(wb["Hazen Output"]["C16"].value, (int, float))
        self.assertIsInstance(
            wb["Geometric accuracy"]["B40"].value, (int, float)
        )
        self.assertTrue((self.tmp / "results.csv").exists())


@pytest.mark.slow
@unittest.skipUnless(
    os.environ.get("HAZEN_GSTT_SESSION"),
    "set HAZEN_GSTT_SESSION to a full local session to run",
)
class TestLocalFullSession(unittest.TestCase):
    """Every step succeeds on a complete local session (not committed)."""

    def test_all_steps_ok(self) -> None:
        protocol = ACRGSTTProtocol(
            os.environ["HAZEN_GSTT_SESSION"],
            report=False,
        )
        outcomes = protocol.run()
        self.assertEqual(
            [(o.step.role, o.step.test) for o in outcomes if o.status != "ok"],
            [],
        )
        self.assertTrue(rows_as_records(protocol.rows(outcomes)))
