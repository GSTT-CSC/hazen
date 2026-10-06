"""Writers for protocol results (Excel template, CSV/TSV and JSON)."""

from __future__ import annotations

# Type Checking
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from collections.abc import Iterable, Mapping, Sequence
    from importlib.resources.abc import Traversable

# Python imports
import csv
import json
import io
import logging
import math
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

# Module imports
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

logger = logging.getLogger(__name__)

Record = dict[str, Any]

# Report images are scaled to this width (px) and laid out side by side
_IMAGE_WIDTH = 360
_IMAGE_GAP = 20
_ROW_HEIGHT = 20  # px, Excel's default row height
# Images are stored at twice their displayed width to stay sharp on zoom
_IMAGE_STORED_WIDTH = 2 * _IMAGE_WIDTH


@dataclass(frozen=True)
class TemplateCell:
    """A value to write into one cell of a workbook template.

    Attributes:
        sheet: Worksheet title.
        ref: Cell reference, e.g. ``"D16"``.
        value: Value to write; None leaves the cell empty.
        note: Comment attached to the cell, e.g. why it has no value.

    """

    sheet: str
    ref: str
    value: Any = None
    note: str = ""


@dataclass(frozen=True)
class ImageGroup:
    """Images shown together under one heading on an image sheet.

    Attributes:
        title: Heading written above the images.
        paths: Image files, laid out left to right and stored as JPEG.
        note: Text written after the heading, e.g. why there are no images.

    """

    title: str
    paths: tuple[str, ...] = ()
    note: str = ""


def fill_template(
    template: str | Path | Traversable,
    cells: Iterable[TemplateCell],
    path: str | Path,
    images: Mapping[str, Sequence[ImageGroup]] | None = None,
) -> Path:
    """Copy a workbook template to *path* with *cells* filled in.

    Formatting and formulas in the template are kept, so formulas
    recalculate when the file is opened.

    Args:
        template: Source ``.xlsx`` workbook, a path or package resource.
        cells: Values (and notes) to write.
        path: Destination ``.xlsx`` file.
        images: Sheets to add after the template's, each a list of
            image groups.

    Returns:
        The path the workbook was written to.

    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    source = Path(template) if isinstance(template, str) else template
    with source.open("rb") as fh:
        wb = load_workbook(fh)
    for cell in cells:
        target = wb[cell.sheet][cell.ref]
        target.value = _cell_value(cell.value)
        if cell.note:
            target.comment = Comment(cell.note, "hazen")
    for title, groups in (images or {}).items():
        _add_image_sheet(wb.create_sheet(title), groups)

    wb.save(path)
    logger.info("Workbook written to %s", path)
    return path


def _add_image_sheet(ws: Any, groups: Sequence[ImageGroup]) -> None:
    """Write each group as a heading row followed by its images."""
    columns = max((len(g.paths) for g in groups), default=0)
    for col in range(1, columns + 1):
        # Excel column width in characters is roughly (px - 5) / 7
        ws.column_dimensions[get_column_letter(col)].width = (
            _IMAGE_WIDTH + _IMAGE_GAP - 5
        ) / 7

    if not groups:
        ws["A1"] = "No report images"
        return

    row = 1
    for group in groups:
        ws.cell(row, 1, group.title).font = Font(bold=True)
        if group.note:
            ws.cell(row, 2, group.note)
        row += 1

        tallest = 0
        for col, image_path in enumerate(group.paths, start=1):
            image = _downscaled(image_path)
            image.height = round(image.height * _IMAGE_WIDTH / image.width)
            image.width = _IMAGE_WIDTH
            tallest = max(tallest, image.height)
            ws.add_image(image, f"{get_column_letter(col)}{row}")
        row += math.ceil(tallest / _ROW_HEIGHT) + 1


def _downscaled(path: str) -> Image:
    """Load an image as a reduced JPEG so the workbook stays small."""
    with PILImage.open(path) as src:
        src.thumbnail((_IMAGE_STORED_WIDTH, src.height))
        flat = PILImage.new("RGB", src.size, "white")
        flat.paste(src, mask=src.getchannel("A") if "A" in src.mode else None)
    buffer = io.BytesIO()
    flat.save(buffer, format="jpeg", quality=85)
    # openpyxl reads the bytes back from the image's file at save time
    return Image(PILImage.open(buffer))


def write_records(
    records: Sequence[Record],
    fmt: str,
    path: str | Path = "-",
) -> None:
    """Write records as json, csv or tsv to *path* (``"-"`` for stdout)."""
    if fmt not in {"json", "csv", "tsv"}:
        msg = f"Unrecognised format {fmt!r}"
        raise ValueError(msg)

    if path == "-":
        _write_records(records, fmt, sys.stdout)
        return
    with Path(path).open("w", newline="") as fh:
        _write_records(records, fmt, fh)


def _write_records(records: Sequence[Record], fmt: str, fh: Any) -> None:
    if fmt == "json":
        json.dump(
            [{k: _cell_value(v) for k, v in r.items()} for r in records],
            fh,
            indent=2,
        )
        fh.write("\n")
        return
    if not records:
        return
    writer = csv.DictWriter(
        fh,
        fieldnames=list(records[0].keys()),
        delimiter="," if fmt == "csv" else "\t",
    )
    writer.writeheader()
    for record in records:
        writer.writerow({k: _cell_value(v) for k, v in record.items()})


def _cell_value(value: Any) -> Any:
    """Convert a value to something Excel/JSON can store."""
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    try:
        return value.item()  # numpy scalar
    except (AttributeError, ValueError):
        return str(value)
